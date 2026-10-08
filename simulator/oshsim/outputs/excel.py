"""פלט Excel עם נוסחאות חיות ו-RTL (סעיף 9א).

היתרה הרצה היא נוסחה, כך שהיועץ יכול לשנות סכום ולראות את ההשפעה. אחרי הכתיבה
`verify_with_libreoffice` מחשב מחדש ב-LibreOffice headless ומשווה לחישוב בקוד.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engine.simulate import SimResult
from ..model.money import to_decimal
from ..model.transaction import Transaction
from ..patterns.recurring import KIND_HE, RecurringPattern

NUM = '#,##0.00;[Red]-#,##0.00'
DATE = 'dd/mm/yyyy'
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
CHANNEL_HE = {"direct": "(י)", "banker": "(פ)", None: ""}
ORIGIN_HE = {"historical": "", "modified": "שונה", "added": "נוסף", "fee": "עמלה מחושבת",
             "interest": "ריבית מחושבת", "projected": "הקרנה"}


def _sheet(wb, title, headers, widths=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _m(d) -> str:
    return f"{d.year:04d}-{d.month:02d}"


def write_excel(path: str | Path, txns: list[Transaction], base: SimResult, scen: SimResult,
                patterns: list[RecurringPattern], param_rows: list[tuple], validation_log: list[str],
                scenario_lines: list[str], credit_limit: int | None) -> dict:
    """→ מפת תאים לבדיקה: {'actual_balance': [(cell, agorot)], 'scenario_balance': [...]}"""
    wb = Workbook()
    wb.remove(wb.active)
    check: dict[str, list] = {"actual_balance": [], "scenario_balance": []}

    summary = _sheet(wb, "סיכום", ["מדד", "בסיס", "תרחיש", "הפרש", "הערה"], [28, 16, 16, 16, 50])

    # ---------- תנועות בפועל ----------
    ws = _sheet(wb, "תנועות בפועל", ["תאריך", "חודש", "סוג תנועה", "ערוץ", "קטגוריה", "תת-קטגוריה",
                                     "אסמכתה", "זכות/חובה", "יתרה רצה", "יתרה בדף", "בדיקה"],
                [12, 9, 34, 6, 16, 16, 10, 14, 14, 14, 10])
    ws.append(["", "", "יתרת פתיחה (נגזרת)", "", "", "", "", None, to_decimal(base.opening)])
    ws.cell(2, 9).number_format = NUM
    for i, t in enumerate(txns, start=3):
        ws.append([t.date, _m(t.date), t.description, CHANNEL_HE.get(t.channel, ""), t.category, t.subcategory,
                   t.reference, to_decimal(t.amount_agorot), f"=I{i - 1}+H{i}",
                   to_decimal(t.day_balance_agorot) if t.day_balance_agorot is not None else None,
                   f'=IF(J{i}="","",ROUND(I{i}-J{i},2))'])
        ws.cell(i, 1).number_format = DATE
        for c in (8, 9, 10):
            ws.cell(i, c).number_format = NUM
        if t.day_balance_agorot is not None:
            check["actual_balance"].append((f"I{i}", t.day_balance_agorot))
    actual_last = len(txns) + 2

    # ---------- תנועות תרחיש ----------
    ws2 = _sheet(wb, "תנועות תרחיש", ["תאריך", "חודש", "סוג תנועה", "ערוץ", "קטגוריה", "מקור", "הערה",
                                      "סכום בסיס", "זכות/חובה", "יתרה רצה"],
                 [12, 9, 34, 6, 16, 12, 40, 14, 14, 14])
    ws2.append(["", "", "יתרת פתיחה", "", "", "", "", None, None, to_decimal(scen.opening)])
    ws2.cell(2, 10).number_format = NUM
    i = 2
    for day in scen.days:
        for r in day.rows:
            i += 1
            ws2.append([r.date, _m(r.date), r.description, CHANNEL_HE.get(r.channel, ""), r.category,
                        ORIGIN_HE.get(r.origin, r.origin), r.note,
                        to_decimal(r.base_amount) if r.base_amount is not None else None,
                        to_decimal(r.amount), f"=J{i - 1}+I{i}"])
            ws2.cell(i, 1).number_format = DATE
            for c in (8, 9, 10):
                ws2.cell(i, c).number_format = NUM
            if r.origin != "historical":
                for c in range(1, 11):
                    ws2.cell(i, c).fill = WARN_FILL
        check["scenario_balance"].append((f"J{i}", day.balance))
    scen_last = i

    # ---------- פער לפי חודש ----------
    months = sorted({_m(d) for d in scen.daily_balance} | {_m(d) for d in base.daily_balance})
    ws3 = _sheet(wb, "פער", ["חודש", "תזרים בסיס", "תזרים תרחיש", "הפרש חודשי", "הפרש מצטבר"], [10, 16, 16, 16, 16])
    for j, m in enumerate(months, start=2):
        ws3.append([m,
                    f"=SUMIFS('תנועות בפועל'!$H$3:$H${actual_last},'תנועות בפועל'!$B$3:$B${actual_last},A{j})",
                    f"=SUMIFS('תנועות תרחיש'!$I$3:$I${scen_last},'תנועות תרחיש'!$B$3:$B${scen_last},A{j})",
                    f"=C{j}-B{j}", f"=D{j}" if j == 2 else f"=E{j - 1}+D{j}"])
        for c in range(2, 6):
            ws3.cell(j, c).number_format = NUM
    gap_last = len(months) + 1

    # ---------- קבועות ודפוסים ----------
    ws4 = _sheet(wb, "קבועות ודפוסים", ["מזהה", "מנפיק", "כיוון", "סוג", "קטגוריה", "יום בחודש", "סכום אופייני",
                                        "כלל שישי", "כלל שבת/חג", "נלמד מהיסטוריה", "אושר", "סכומים לפי חודש"],
                 [7, 30, 7, 9, 16, 9, 14, 9, 10, 10, 6, 60])
    for p in patterns:
        ws4.append([p.id, p.counterparty, "זכות" if p.direction == "credit" else "חובה", KIND_HE[p.kind], p.category,
                    p.day_of_month, to_decimal(p.typical_amount), p.friday_rule, p.closed_rule,
                    "כן" if p.rule_learned else "ברירת מחדל", "כן" if p.approved else "לא",
                    "; ".join(f"{m}: {to_decimal(a)}" for m, a in p.amounts_by_month.items())])
        ws4.cell(ws4.max_row, 7).number_format = NUM

    # ---------- קטגוריות לפי חודש ----------
    cats = sorted({t.category for t in txns} | {r.category for r in scen.rows() if r.category})
    ws5 = _sheet(wb, "קטגוריות לפי חודש", ["קטגוריה", "גרסה"] + months, [18, 8] + [13] * len(months))
    row = 2
    for cat in cats:
        for label, sheet, amt_col, last in (("בסיס", "תנועות בפועל", "H", actual_last),
                                            ("תרחיש", "תנועות תרחיש", "I", scen_last)):
            cat_col = "E"
            ws5.append([cat, label] + [
                f"=SUMIFS('{sheet}'!${amt_col}$3:${amt_col}${last},'{sheet}'!${cat_col}$3:${cat_col}${last},$A{row},"
                f"'{sheet}'!$B$3:$B${last},{get_column_letter(k + 3)}$1)" for k in range(len(months))])
            for k in range(len(months)):
                ws5.cell(row, k + 3).number_format = NUM
            row += 1

    # ---------- הנחות ופרמטרים ----------
    ws6 = _sheet(wb, "הנחות ופרמטרים", ["פרמטר", "ערך", "מקור", "תאריך/הערה"], [40, 30, 50, 30])
    for r in param_rows:
        ws6.append(list(r))
    ws6.append([])
    ws6.append(["תרחיש", scen.name])
    for line in scenario_lines:
        ws6.append(["שינוי", line])
    ws6.append([])
    ws6.append(["כיול עמלת ערוץ ישיר", "חודש", "פעולות (י) בדף / מונה באסמכתה", "מודל מול בפועל"])
    for c in base.fee_calibration:
        ws6.append(["", c.month, f"{c.observed_count} / {c.reference_count if c.reference_count is not None else '—'}",
                    "—" if c.model_amount is None else
                    f"{to_decimal(c.model_amount)} מול {to_decimal(c.actual_amount)} "
                    f"{'✓' if c.ok else '✗'}"])
    ws6.append([])
    ws6.append(["כיול ריבית", "תקופה", "מודל מול בפועל", "טעות"])
    if not base.interest_calibration:
        ws6.append(["", "לא בוצע — לא הוזן שיעור ריבית", "", ""])
    for c in base.interest_calibration:
        ws6.append(["", f"{c.period_start:%d/%m/%Y}–{c.period_end:%d/%m/%Y}"
                    + (" (תקופה חלקית בדף)" if c.partial_period else ""),
                    f"{to_decimal(c.model_amount)} מול {to_decimal(c.actual_amount)}",
                    f"{c.error_pct}%" if c.error_pct is not None else "—"])
    ws6.append([])
    for n in scen.notes:
        ws6.append(["הערה", n])

    # ---------- יומן ולידציה ----------
    ws7 = _sheet(wb, "יומן ולידציה", ["רשומה"], [120])
    for line in validation_log:
        ws7.append([line])

    # ---------- סיכום (נוסחאות מול הגיליונות) ----------
    def put(label, b, s, note="", fmt_num=True):
        summary.append([label, b, s, f"=C{summary.max_row + 1}-B{summary.max_row + 1}" if fmt_num else "", note])
        for c in (2, 3, 4):
            summary.cell(summary.max_row, c).number_format = NUM if fmt_num else "0"

    put("יתרת פתיחה", f"='תנועות בפועל'!I2", f"='תנועות תרחיש'!J2")
    put("יתרת סגירה", f"='תנועות בפועל'!I{actual_last}", f"='תנועות תרחיש'!J{scen_last}",
        "תרחיש Forward — כולל הקרנה (הערכה)" if scen.projected_from else "")
    put("יתרה מינימלית (ימי תנועה)", f"=MIN('תנועות בפועל'!I3:I{actual_last})",
        f"=MIN('תנועות תרחיש'!J3:J{scen_last})")
    put("הפרש מצטבר", None, f"='פער'!E{gap_last}")
    put("סך עמלות ערוץ ישיר", to_decimal(base.total_of("direct_channel_fee")),
        to_decimal(scen.total_of("direct_channel_fee")), "מחושב בקוד")
    put("סך ריבית", to_decimal(base.total_of("interest")), to_decimal(scen.total_of("interest")),
        "הערכה" if scen.interest_is_estimate else "מכויל")
    put("ימים בחריגה ממסגרת", base.days_over_limit(credit_limit), scen.days_over_limit(credit_limit),
        "מסגרת לא הוזנה" if credit_limit is None else "כל ימי הלוח, מחושב בקוד", fmt_num=False)
    summary.append([])
    summary.append(["סימולציה — אינו מסמך בנקאי"])
    summary.cell(summary.max_row, 1).font = Font(bold=True, color="C00000", size=14)

    wb.save(str(path))
    return check


def _soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def verify_with_libreoffice(path: str | Path, check: dict) -> list[str]:
    """מחשב מחדש ב-LibreOffice ומשווה את היתרות הרצות לחישוב בקוד. → רשימת בעיות."""
    exe = _soffice()
    if not exe:
        return ["LibreOffice לא מותקן — לא בוצע חישוב מחדש"]
    path = Path(path)
    problems: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                       check=True, capture_output=True, timeout=180)
        wb = load_workbook(Path(tmp) / path.name, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:")):
                        problems.append(f"{ws.title}!{c.coordinate}: {c.value}")
        sheets = {"actual_balance": "תנועות בפועל", "scenario_balance": "תנועות תרחיש"}
        for key, sheet in sheets.items():
            ws = wb[sheet]
            for cell, agorot in check[key]:
                v = ws[cell].value
                if v is None:
                    problems.append(f"{sheet}!{cell}: אין ערך מחושב")
                    continue
                got = Decimal(str(v)).quantize(Decimal("0.01"))
                if got != to_decimal(agorot):
                    problems.append(f"{sheet}!{cell}: Excel {got} מול קוד {to_decimal(agorot)}")
    return problems
