import re
from decimal import Decimal

import pytest

from oshsim.engine.deltas import ChangeAmount, Selector
from oshsim.engine.simulate import Scenario
from oshsim.outputs.excel import _soffice, verify_with_libreoffice, write_excel
from oshsim.outputs.html import WATERMARK, chromium_path, html_to_pdf, report_html, statement_html
from oshsim.pipeline import run_scenario
from oshsim.privacy.masking import mask_account, mask_text
from oshsim.privacy.workspace import Workspace, generate_key


@pytest.fixture
def run_out(loaded, params_rates):
    sc = Scenario("שכר +10%", deltas=[ChangeAmount(Selector(category="הכנסות", subcategory="משכורת"),
                                                  pct=Decimal(10), label="שכר +10%")])
    return run_scenario(loaded.txns, loaded.report.opening_agorot, params_rates, sc)


@pytest.mark.skipif(not _soffice(), reason="LibreOffice לא מותקן")
def test_excel_recalc_matches_code(tmp_path, loaded, run_out, params_rates):
    p = tmp_path / "out.xlsx"
    check = write_excel(p, loaded.txns, run_out.base, run_out.scenario, run_out.engine.patterns,
                        params_rates.table_rows(), loaded.log, run_out.scenario_lines, params_rates.credit_limit_agorot)
    assert check["actual_balance"] and check["scenario_balance"]
    assert verify_with_libreoffice(p, check) == []
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert wb.sheetnames == ["סיכום", "תנועות בפועל", "תנועות תרחיש", "פער", "קבועות ודפוסים",
                             "קטגוריות לפי חודש", "הנחות ופרמטרים", "יומן ולידציה"]
    assert all(ws.sheet_view.rightToLeft for ws in wb.worksheets)
    assert str(wb["תנועות בפועל"]["I3"].value).startswith("=")      # יתרה רצה כנוסחה


def test_statement_html_marking(run_out):
    h = statement_html(run_out.scenario, account_hint="12-345-678901")
    assert h.count(WATERMARK) >= 3                                   # סימן מים, באנר, כותרת טבלה חוזרת
    assert 'class="wm"' in h and "position:fixed" in h and "table-header-group" in h
    assert "678901" not in h and "8901" in h                         # מספר חשבון מוסתר חלקית
    assert "מזרחי" not in h and "טפחות" not in re.sub(r"<td>[^<]*טפחות[^<]*</td>", "", h)   # בלי שם הבנק כמנפיק
    assert 'dir="rtl"' in h


def test_report_html(run_out):
    h = report_html(run_out.base, run_out.scenario, 4000000, run_out.scenario_lines)
    assert WATERMARK in h and "<svg" in h and "השפעות משניות" in h and "שכר +10%" in h


@pytest.mark.skipif(not chromium_path(), reason="Chromium לא זמין")
def test_pdf_every_page_marked(tmp_path, run_out):
    import pdfplumber

    from oshsim.ingest.hebrew import visual_to_logical
    hp = tmp_path / "s.html"
    hp.write_text(statement_html(run_out.scenario), encoding="utf-8")
    pp = tmp_path / "s.pdf"
    assert html_to_pdf(hp, pp)
    with pdfplumber.open(pp) as pdf:
        assert len(pdf.pages) > 1
        for page in pdf.pages:
            t = page.extract_text() or ""
            assert "סימולציה" in t or "סימולציה" in visual_to_logical(t)


def test_masking():
    assert mask_account("12-345-678901") == "•" * 7 + "8901"
    assert mask_text("כרטיס 4580 1234 5678 9012 חיוב") == "כרטיס " + "•" * 12 + "9012 חיוב"


def test_workspace_encryption_and_purge(tmp_path, monkeypatch):
    monkeypatch.setenv("OSHSIM_KEY", generate_key())
    ws = Workspace(tmp_path / "c")
    ws.write_json("ledger.json", {"x": "סוד"})
    raw = (tmp_path / "c" / "ledger.json.enc").read_bytes()
    assert "סוד".encode() not in raw and not (tmp_path / "c" / "ledger.json").exists()
    assert ws.read_json("ledger.json") == {"x": "סוד"}
    log = (tmp_path / "c" / "access.log").read_text(encoding="utf-8")
    assert '"write"' in log and '"read"' in log
    monkeypatch.delenv("OSHSIM_KEY")
    with pytest.raises(PermissionError):
        Workspace(tmp_path / "c").read_json("ledger.json")
    assert Workspace(tmp_path / "c").purge() >= 2
    assert not (tmp_path / "c").exists()


def test_compare_rows_and_page(loaded, run_out):
    from oshsim.outputs.compare import compare_html, compare_rows
    rows = compare_rows(loaded.txns, loaded.report.opening_agorot, run_out.scenario)
    assert len(rows) == len(loaded.txns)
    changed = [r for r in rows if r.status == "changed"]
    # המשכורות שונו, ושורת הריבית הותאמה ליתרה הגבוהה יותר
    assert changed and all("משכורת" in r.description or "ריבית" in r.description for r in changed)
    assert any("משכורת" in r.description for r in changed)
    # יתרה רצה בכל צד, ובסוף כל יום — יתרת הדף / יתרת התרחיש
    for r, t in zip(rows, loaded.txns):
        if t.day_balance_agorot is not None:
            assert r.day_end and r.old_balance == t.day_balance_agorot
            assert r.new_balance == run_out.scenario.daily_balance[t.date]
    assert rows[-1].balance_gap == run_out.scenario.closing - run_out.base.closing
    h = compare_html(loaded.txns, loaded.report.opening_agorot, run_out.scenario, "12-345-678901")
    assert "דף אמיתי" in h and "דף חדש" in h and WATERMARK in h and "678901" not in h
    assert h.count('class="changed') == len(changed)


def test_replica_marks_every_change(loaded, run_out):
    from oshsim.outputs.compare import compare_rows
    from oshsim.outputs.replica import replica_html
    h = replica_html(loaded.txns, loaded.report.opening_agorot, run_out.scenario, "511-237920", "תקופה לבדיקה")
    rows = compare_rows(loaded.txns, loaded.report.opening_agorot, run_out.scenario)
    changed = sum(r.status == "changed" for r in rows)
    bal_changed = sum(1 for r in rows if r.day_end and r.new_balance != r.old_balance)
    assert h.count('class="amt chg"') == changed and h.count('class="bal balchg"') == bal_changed
    assert WATERMARK in h and "237920" not in h and "7920" in h and "תקופה לבדיקה" in h
    assert "מזרחי" not in h and "טפחות" not in h.split("<tbody>")[0]   # בלי שם הבנק בכותרת
